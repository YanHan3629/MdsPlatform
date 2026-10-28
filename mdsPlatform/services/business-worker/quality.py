"""Measured, read-only snapshot quality. Uses the existing format parsers, no LLM.

Rules are source declarations, never guessed business keys. Optional declarations
come from catalog/field_dictionary.json, entities.json, relations.json and
assets.json, or an explicit per-path rule configuration captured in the report.
"""
from __future__ import annotations
import hashlib
import json
import re
import csv
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path, PurePosixPath
from zoneinfo import ZoneInfo
from professional.contracts import ParseContext
from professional.parsers import default_parser_registry

RULES_VERSION = "snapshot-quality-2026-10-20.6"
IMPLEMENTATION_DIGEST=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
RULE_NAMES = ("schema", "primary_key", "references", "units", "time", "nulls")
UNITS = {"1", "%", "mm", "cm", "m", "m2", "m²", "m3", "m³", "L", "l", "ml", "kg", "g", "t", "s", "ms", "min", "h", "d", "day", "Hz", "kHz", "rpm", "Pa", "kPa", "MPa", "bar", "N", "N·m", "Nm", "J", "kJ", "W", "kW", "Wh", "kWh", "V", "A", "mA", "Ω", "ohm", "℃", "°C", "C", "K", "件", "台", "个", "套", "元", "CNY", "RMB", "万元", "元/件", "元/台", "人", "次", "无量纲"}
UNITS.update({"km","mm/s","mL","箱","mm2","mm3","kΩ","通过计数","片"})
csv.field_size_limit(16*1024*1024)

def canonical_manifest(files):
    return hashlib.sha256(json.dumps(sorted(files, key=lambda f: f["logicalPath"]), ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()

def safe_path(path):
    parts = PurePosixPath(path).parts
    if not path.startswith("/") or path.startswith("//") or "\\" in path or "\x00" in path or any(p in ("..", ".") for p in path.split("/")[1:]) or "//" in path:
        raise ValueError("unsafe logical path")
    return Path(*parts[1:])

def check(path, rule, outcomes=None, *, state=None, critical=False, evidence=None):
    outcomes = outcomes or []
    failures = [e for ok, e in outcomes if not ok]
    return dict(logicalPath=path, rule=rule, status=state or ("FAIL" if failures else "PASS"), critical=critical,
                denominator=len(outcomes), passed=len(outcomes)-len(failures), issueCount=len(failures),
                issues=failures[:100], evidence=evidence or {}, truncated=len(failures)>100)

def null(value):
    return value is None or (isinstance(value, str) and not value.strip())

def value_type(value, declared):
    if "|" in declared:
        results=[value_type(value,t.strip()) for t in declared.split("|")]
        return True if True in results else None if None in results else False
    if null(value):
        return True
    if declared in ("string", "str", "text"):
        return isinstance(value, str)
    if declared in ("int", "integer", "int64", "number", "float", "float64", "double", "numeric"):
        if isinstance(value, bool): return False
        try:
            x=float(value)
            return x == x and abs(x) != float("inf") and (declared not in ("int", "integer", "int64") or x.is_integer())
        except (TypeError, ValueError): return False
    if declared in ("bool", "boolean"):
        return isinstance(value, bool) or str(value).lower() in ("true", "false", "0", "1")
    if declared in ("array", "list"): return isinstance(value,list)
    if declared in ("object", "dict"): return isinstance(value,dict)
    if declared=="null":return null(value)
    if declared.startswith(("date","datetime","timestamp")) or "excel日期" in declared:return True # separately evaluated with time rule
    return None

def timezone_valid(timezone):
    """Validate an explicit declaration independently of the cell's encoding."""
    if timezone is None or timezone == "":
        return None
    if not isinstance(timezone, str):
        return False
    offset = re.fullmatch(r"[+-](\d{2}):(\d{2})", timezone)
    if offset:
        return int(offset[1]) < 24 and int(offset[2]) < 60
    if timezone == "Z":
        return True
    try:
        ZoneInfo(timezone)
        return True
    except (ValueError, KeyError):
        return False

def time_valid(value, granularity, timezone, declared=""):
    zone_valid = timezone_valid(timezone)
    if zone_valid is False: return False
    if null(value): return True
    text=str(value).strip()
    if "excel" in declared.lower() and isinstance(value,(int,float)):
        if isinstance(value, bool): return False
        try:datetime(1899,12,30)+timedelta(days=float(value));return zone_valid
        except (ValueError,OverflowError):return False
    if granularity in ("date", "day", "daily", "日期", "日") or (re.fullmatch(r"\d{4}-\d{2}-\d{2}",text) and not granularity):
        try: datetime.strptime(text,"%Y-%m-%d");return True
        except ValueError:return False
    try:
        parsed=datetime.fromisoformat(text.replace("Z","+00:00"))
        if parsed.tzinfo is None:
            return zone_valid
        return True
    except (ValueError, KeyError):return False

def json_field_values(document,field):
    """Resolve declared nested paths against the original JSON, retaining locators."""
    tokens=field.split(".")
    def walk(value,rest,location):
        if not rest:return [(dict(jsonPath=location),value,True,{})]
        # A root JSON list is a table; unprefixed fields apply to each record.
        if isinstance(value,list):
            return [result for i,child in enumerate(value) for result in walk(child,rest,f"{location}[{i}]")]
        token=rest[0];array=token.endswith("[]");key=token[:-2] if array else token
        if not isinstance(value,dict) or key not in value:return [(dict(jsonPath=location,field=field),None,False,{})]
        child=value[key];here=f"{location}.{key}"
        if array:
            if not isinstance(child,list):return [(dict(jsonPath=here),child,False,value)]
            return [result for i,item in enumerate(child) for result in walk(item,rest[1:],f"{here}[{i}]")]
        if len(rest)==1:return [(dict(jsonPath=here),child,True,value)]
        return walk(child,rest[1:],here)
    return walk(document,tokens,"$")

def unit_result(unit,values,config):
    if not isinstance(unit,str):return False,str(unit)[:100]
    custom=config.get("additionalUnits") or {}
    if unit in UNITS or isinstance(custom,dict) and isinstance(custom.get(unit),str) and custom[unit].strip():return True,unit
    if unit in ("不适用","不适用（单位名称）","单位枚举"):return "NA",unit
    if unit in ("CNY含税","无量纲，0/1"):return True,unit
    if "见" in str(unit) or "由记录" in str(unit):
        candidates=re.findall(r"quantity_checked_unit|quantity_unit|inspection_unit|measurement_unit|unit",str(unit))
        for key in candidates:
            if key in values and not null(values[key]):
                actual=str(values[key]);return actual in UNITS or isinstance(custom,dict) and isinstance(custom.get(actual),str) and bool(custom[actual].strip()),actual
        return None,unit
    return False,unit

def observed_json_fields(value,prefix="",result=None):
    result=set() if result is None else result
    if isinstance(value,dict):
        for key,child in value.items():
            name=prefix+"."+key if prefix else key;result.add(name);observed_json_fields(child,name,result)
    elif isinstance(value,list):
        for child in value:observed_json_fields(child,prefix+"[]",result)
    return result

def inspect_snapshot(job, root, download, progress=lambda:None):
    files=job["files"]; checks=[]; paths={f["logicalPath"] for f in files}; valid={}; parsed={}; docs={}; tables={};formula_cells={}
    identities=canonical_manifest(files)
    checks.append(check("/", "manifest_digest", [(identities==job["manifestDigest"], {"expected":job["manifestDigest"],"actual":identities})],critical=True))
    registry=default_parser_registry()
    for f in files:
        progress(); path=f["logicalPath"]
        try:
            local=root/safe_path(path);local.parent.mkdir(parents=True,exist_ok=True)
            download(path,local)
            digest=hashlib.sha256()
            with local.open("rb") as stream:
                for block in iter(lambda:stream.read(1024*1024),b""):digest.update(block)
            size=local.stat().st_size
            ok=digest.hexdigest()==f["sha256"] and size==f["sizeBytes"]
            identity={"expectedSha256":f["sha256"],"actualSha256":digest.hexdigest(),"expectedSize":f["sizeBytes"],"actualSize":size}
            checks.append(check(path,"file_integrity",[(ok,identity)],critical=True,evidence=identity))
            if not ok:continue
            valid[path]=local
        except Exception as exc:
            checks.append(check(path,"file_integrity",[(False,{"error":type(exc).__name__,"reason":"文件无法读取或路径无效"})],critical=True));continue
        parser=registry.resolve(local)
        if parser is None:
            checks.append(check(path,"format_parse",state="NOT_ASSESSED",critical=True,evidence={"extension":local.suffix}));continue
        try:
            asset=parser.parse(local,ParseContext(job["versionId"],path,root.parent/"derived"))
            if local.suffix.lower() in (".png",".jpg",".jpeg"):
                from PIL import Image
                with Image.open(local) as image:image.verify()
                with Image.open(local) as image:image.load()
            parsed[path]=asset
            tables[path]=[(item.locator,item.payload["records"]) for item in asset.items if item.kind=="table"]
            if local.suffix.lower()==".csv":
                # CSV stores lexical cells. Parser dtype inference must not turn a
                # declared string 'True' into a schema error, or lose leading zeroes.
                with local.open(encoding="utf-8-sig",newline="") as stream:
                    rows=[dict(rawRow=i,values=dict(row)) for i,row in enumerate(csv.DictReader(stream),2)]
                tables[path]=[(dict(sheet="root",rawRowStart=2),rows)]
            if local.suffix.lower()==".xlsx":
                from openpyxl import load_workbook
                formulas={};workbook=load_workbook(local,read_only=True,data_only=False)
                try:
                    for item in asset.items:
                        if item.kind!="table":continue
                        sheet=item.locator["sheet"];columns=[c["name"] for c in item.payload["columns"]];start=item.locator["rawRowStart"]
                        for row in workbook[sheet].iter_rows(min_row=start,max_col=len(columns)):
                            for column,cell in zip(columns,row):
                                if cell.data_type=="f":formulas[(sheet,cell.row,column)]=str(cell.value)
                    formula_cells[path]=formulas
                finally:workbook.close()
            if local.suffix.lower()==".json":docs[path]=json.loads(local.read_text(encoding="utf-8-sig"))
            checks.append(check(path,"format_parse",[(True,{})],critical=True,evidence={"parser":asset.parser_name,"parserVersion":asset.parser_version,"itemCount":len(asset.items),"qualityFlags":asset.quality_flags}))
        except Exception as exc:
            checks.append(check(path,"format_parse",[(False,{"error":type(exc).__name__,"reason":"真实格式解析失败"})],critical=True))
    # No source declaration means unassessed, not inferred success.
    declarations=job.get("rules",{}).get("files",{})
    declared_fields=defaultdict(list); pk_by_path=defaultdict(set)
    for field in docs.get("/catalog/field_dictionary.json",[]):
        if isinstance(field,dict) and isinstance(field.get("path"),str) and isinstance(field.get("field"),str) and field["path"] and field["field"]:
            declared_fields["/"+field["path"].lstrip("/")].append(field)
    # These are metadata-format contracts, not guessed domain/business schemas.
    catalog_schemas={
        "/catalog/assets.csv":{"asset_id":"string","path":"string","sha256":"string"},
        "/catalog/field_dictionary.json":{"path":"string","field":"string","type":"string"},
        "/catalog/assets.json":{"asset_id":"string","path":"string","sha256":"string"},
        "/catalog/entities.json":{"entity_id":"string","path":"string"},
        "/catalog/relations.json":{"source_id":"string","target_id":"string","path":"string"},
        "/catalog/retrieval_manifest.json":{"version":"string","path_base":"string","scope":"string","files":"array"},
        "/catalog/sources_and_assumptions.json":{"sources":"array","assumptions":"array","simulation_notice":"string","dataset_version":"string","version":"string"},
    }
    for path,fields in catalog_schemas.items():
        if path in paths:declared_fields[path]=[dict(field=k,type=v,nullable=False) for k,v in fields.items()]
    entities=docs.get("/catalog/entities.json",[]); relations=docs.get("/catalog/relations.json",[])
    if not isinstance(entities,list):entities=[]
    if not isinstance(relations,list):relations=[]
    entity_ids={str(e["entity_id"]) for e in entities if isinstance(e,dict) and e.get("entity_id") is not None}
    for e in entities:
        if isinstance(e,dict) and isinstance(e.get("primary_key"),str) and isinstance(e.get("path"),str) and e["primary_key"] and e["path"] and e.get("identity_kind")!="asset_catalog" and re.match(r"^row \d+ "+re.escape(e["primary_key"])+"=",str(e.get("locator",""))):
            pk_by_path["/"+e["path"].lstrip("/")].add(e["primary_key"])
    # Resolve referenced tables only within this verified snapshot.
    def rows_for(path,sheet=None):
        return [(locator, row) for locator,rows in tables.get(path,[]) if sheet is None or locator.get("sheet")==sheet for row in rows]
    all_issues_by_rule=Counter()
    for path,asset in parsed.items():
        progress(); groups=tables[path]; config=declarations.get(path,declarations.get(path.lstrip("/"),{}))
        supplied=config.get("fields",[])
        if isinstance(supplied,dict):supplied=[dict(v,field=k) for k,v in supplied.items()]
        fields=list(declared_fields[path]);source_fields={(f.get("sheet"),f["field"]):f for f in fields}
        for field in supplied:
            key=(field.get("sheet"),field["field"]);original=source_fields.get(key)
            if original is not None:
                conflicts=[name for name in ("type","nullable","unit","timezone","null_rule","null_values") if name in field and name in original and field[name]!=original[name]]
                if conflicts:checks.append(check(path,"published_declaration_conflict",[(False,{"field":field["field"],"conflictingConstraints":conflicts,"reason":"本次规则不能覆盖已发布源字典的约束；需显式修订源版本"})],critical=True))
            else:fields.append(field)
        if not groups and path not in docs:
            for rule in RULE_NAMES:checks.append(check(path,rule,state="NOT_ASSESSED" if rule=="schema" and supplied else "NOT_APPLICABLE",critical=bool(rule=="schema" and supplied),evidence={"reason":"该解析表示不含表格行；明确请求的字段检查无法执行，其余表格规则不适用"}))
            if path.lower().endswith((".obj",".step",".stp")):
                source_unit=asset.attributes.get("unit") or next((item.units.get("length") for item in asset.items if item.kind=="geometry" and item.units.get("length")),None)
                if source_unit and config.get("geometryUnit") and source_unit!=config["geometryUnit"]:checks.append(check(path,"published_geometry_unit_conflict",[(False,{"sourceUnit":source_unit,"requestedUnit":config["geometryUnit"],"reason":"质量规则不能替换原文件声明的几何单位"})],critical=True))
                unit=source_unit or config.get("geometryUnit")
                checks.append(check(path,"geometry_units",[(unit in UNITS,{"unit":unit})] if unit else None,state=None if unit else "NOT_ASSESSED",evidence={"unit":unit,"reason":"几何单位必须来自原文件或明确映射"}))
            continue
        if not fields:
            checks.append(check(path,"schema",state="NOT_ASSESSED",critical=path.lower().endswith((".csv",".xlsx",".json",".jsonl")),evidence={"reason":"表格没有声明字段类型；请补充规则或源字段字典"}))
            for rule in ("units","time","nulls"):checks.append(check(path,rule,state="NOT_ASSESSED",evidence={"reason":"无源字段声明"}))
        else:
            outcomes={name:[] for name in ("schema","units","time","nulls")};unassessed=Counter();formula_results_unassessed=0
            for field in fields:
                name=field.get("field"); sheet=field.get("sheet")
                if path.lower().endswith(".xlsx") and "." in str(name):
                    qualifier,column=name.split(".",1)
                    if qualifier in {locator.get("sheet") for locator,_ in groups} and (sheet is None or sheet==qualifier):sheet=qualifier;name=column
                if path not in docs:
                    applicable=[item for item in asset.items if item.kind=="table" and (sheet is None or item.locator.get("sheet")==sheet)]
                    if not applicable or not any(name in {column["name"] for column in item.payload["columns"]} for item in applicable):outcomes["schema"].append((False,{"field":name,"sheet":sheet,"reason":"声明的工作表或字段未在实际表头中出现"}))
                selected=[(dict(locator,row=row["rawRow"],field=name),row["values"].get(name),name in row["values"],row["values"]) for locator,row in rows_for(path,sheet)]
                if path in docs:selected=json_field_values(docs[path],name)
                for where,value,present,values in selected:
                    if not present:
                        outcomes["schema"].append((False,dict(where,reason="声明字段不存在")));continue
                    declared=str(field.get("type","")).lower(); formula=formula_cells.get(path,{}).get((where.get("sheet"),where.get("row"),name))
                    schema_value=value
                    if formula is not None and declared in ("string","str","text") and ("公式" in str(field.get("calculation","")) or field.get("formulaType")=="expression"):
                        schema_value=formula;formula_results_unassessed+=1
                    result=value_type(schema_value,declared)
                    if result is None:unassessed["schema"]+=1
                    else:outcomes["schema"].append((result,dict(where,declaredType=field.get("type"),value=str(value)[:100])))
                    nullable=field.get("nullable")
                    if nullable is not None and not isinstance(nullable,bool):
                        outcomes["nulls"].append((False,dict(where,reason="源规则nullable必须是真正的布尔值")));nullable=None
                    if nullable is None:unassessed["nulls"]+=1
                    else:outcomes["nulls"].append((bool(nullable) or not null(value),dict(where,reason="非空字段为空")))
                    gran=field.get("time_granularity") or field.get("timeGranularity")
                    if declared.startswith(("datetime","timestamp","date")) or "excel日期" in declared:
                        timezone=field.get("timezone") or config.get("timezone")
                        if not timezone and "Asia/Shanghai" in str(gran)+declared:timezone="Asia/Shanghai"
                        if not timezone:
                            match=re.search(r"[+-]\d{2}:\d{2}",str(gran)+declared)
                            if match:timezone=match[0]
                        t=time_valid(value,gran,timezone,declared)
                        if t is None:unassessed["time"]+=1
                        else:outcomes["time"].append((t,dict(where,value=str(value)[:100],reason="时间无法解析或时区无效")))
                    if field.get("unit") is not None:
                        unit_ok,actual=unit_result(field["unit"],values,config)
                        if unit_ok is None:unassessed["units"]+=1
                        elif unit_ok!="NA":outcomes["units"].append((unit_ok,dict(where,unit=field["unit"],resolvedUnit=actual,reason="按源字段单位或本行单位列核对")))
                # Nulls with nullable=true still require explanatory semantics.
                if field.get("nullable") is True:
                    outcomes["nulls"].append((bool(field.get("null_rule") or field.get("nullRule")),{"field":name,"reason":"允许空值需声明含义，不能把空值当零"}))
            for rule,results in outcomes.items():
                checks.append(check(path,rule,results,state=None if results else "NOT_ASSESSED" if unassessed[rule] else "NOT_APPLICABLE",critical=rule in ("schema","time"),evidence={"unknownAssertions":unassessed[rule]}))
                if unassessed[rule] and results:checks.append(check(path,rule+"_unassessed",state="NOT_ASSESSED",critical=rule in ("schema","time"),evidence={"assertions":unassessed[rule]}))
            if formula_results_unassessed:checks.append(check(path,"formula_result_type",state="NOT_ASSESSED",evidence={"assertions":formula_results_unassessed,"reason":"源字典声明公式表达式为字符串；表达式已检查，缓存结果类型没有独立声明，不把数值缓存误报为表达式类型错误"}))
            if path in docs:
                observed=observed_json_fields(docs[path]);defined={f["field"] for f in fields};unknown=sorted(observed-defined)
            else:
                observed=set();defined=set()
                for item in asset.items:
                    if item.kind!="table":continue
                    sheet=item.locator.get("sheet");observed.update((str(sheet),c["name"]) for c in item.payload["columns"])
                    for field in fields:
                        name=field["field"];qualifier=field.get("sheet")
                        if path.lower().endswith(".xlsx") and "." in name:
                            prefix,column=name.split(".",1)
                            if prefix==sheet and (qualifier is None or qualifier==sheet):name=column;qualifier=sheet
                        if qualifier is None or qualifier==sheet:defined.add((str(sheet),name))
                unknown=[dict(sheet=s,column=c) for s,c in sorted(observed-defined)]
            if unknown:checks.append(check(path,"schema_coverage",state="NOT_ASSESSED",evidence={"observedFieldCount":len(observed),"unassessedFieldCount":len(unknown),"unassessedFields":unknown[:100],"truncated":len(unknown)>100,"reason":"未声明字段不算已检查；通过率仅针对已测断言"}))
        keys=config.get("primaryKey")
        if isinstance(keys,str):keys=[keys]
        if len(pk_by_path[path])==1:
            published_keys=list(pk_by_path[path])
            if keys is not None and keys!=published_keys:checks.append(check(path,"published_primary_key_conflict",[(False,{"sourcePrimaryKey":published_keys,"requestedPrimaryKey":keys,"reason":"不能用本次规则跳过已发布的主键声明"})],critical=True))
            keys=published_keys
        if keys:
            results=[]
            for locator,rows in groups:
                seen={}
                for row in rows:
                    values=row["values"]; key=tuple(str(values.get(k)) for k in keys)
                    present=all(k in values and not null(values.get(k)) for k in keys)
                    results.append((present and key not in seen,dict(locator,row=row["rawRow"],fields=keys,firstRow=seen.get(key),value=list(key))))
                    seen.setdefault(key,row["rawRow"])
            checks.append(check(path,"primary_key",results,critical=True))
        else:checks.append(check(path,"primary_key",state="NOT_ASSESSED",evidence={"reason":"未声明主键；不以相似列名猜测主键"}))
        references=config.get("references") or []; results=[]
        for ref in references:
            target="/"+ref["path"].lstrip("/"); target_values={str(r["values"].get(ref["targetField"])) for _,r in rows_for(target,ref.get("targetSheet"))}
            for locator,row in rows_for(path,ref.get("sheet")):
                value=row["values"].get(ref["field"])
                results.append((target in paths and (null(value) and ref.get("nullable",False) or str(value) in target_values),dict(locator,row=row["rawRow"],field=ref["field"],value=str(value)[:100],targetPath=target)))
        for relation in relations:
            if isinstance(relation,dict) and "/"+str(relation.get("path","")).lstrip("/")==path:
                results.append((str(relation.get("source_id")) in entity_ids and str(relation.get("target_id")) in entity_ids,{"locator":relation.get("locator"),"field":relation.get("field"),"sourceId":relation.get("source_id"),"targetId":relation.get("target_id")}))
        checks.append(check(path,"references",results,state=None if results else "NOT_ASSESSED",critical=bool(results),evidence={"reason":"引用按声明表或目录实体核对，不推测跨格式关联"}))
    # Catalog path/hash and locator references are measured independently.
    assets=docs.get("/catalog/assets.json",[])
    if isinstance(assets,list) and assets:
        manifest={f["logicalPath"]:f for f in files};results=[]
        for a in assets:
            if not isinstance(a,dict):continue
            p="/"+str(a.get("path","")).lstrip("/")
            results.append((p in manifest and a.get("sha256")==manifest[p]["sha256"],{"catalogAsset":a.get("asset_id"),"targetPath":p,"reason":"目录引用及哈希与冻结清单一致"}))
        checks.append(check("/catalog/assets.json","catalog_references",results,critical=True))
    by_hash=defaultdict(list)
    for f in files:by_hash[f["sha256"]].append(f["logicalPath"])
    same_content=[p for p in by_hash.values() if len(p)>1]
    # Catalog entity identities across carrier formats are reports, never deletion decisions.
    carriers=defaultdict(set)
    for e in entities:
        if isinstance(e,dict) and e.get("entity_id") and e.get("path"):carriers[str(e["entity_id"])].add(e["path"])
    cross=[dict(entityId=k,paths=sorted(v)) for k,v in carriers.items() if len(v)>1]
    return dict(rulesVersion=RULES_VERSION,manifestDigest=job["manifestDigest"],checks=checks,coverage=dict(fileCount=len(files),verifiedFileCount=len(valid),parsedFileCount=len(parsed),
        identicalHashGroups=same_content,crossFormatEntityGroupCount=len(cross),crossFormatEntityGroups=cross[:100],
        ruleImplementationSha256=IMPLEMENTATION_DIGEST,
        note="跨格式实体可能为同一业务事实；仅报告来源关系，不改写、删除或把载体数当业务记录数"))
