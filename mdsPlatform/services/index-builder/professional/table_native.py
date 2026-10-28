"""Source-preserving CSV and OOXML reads. Never execute spreadsheet formulas."""
from __future__ import annotations

import csv
import io
import math
import posixpath
import re
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from zipfile import ZipFile

from .contracts import ParsedItem
from .json_native import records_payload

RECIPE = "native-table-2026-10-24.3"
# Real source catalog cells exceed the stdlib's 128 KiB default. Keep a
# documented bound and reject overflow; never shorten an original field.
CSV_MAX_FIELD_CHARACTERS = 16 * 1024 * 1024
csv.field_size_limit(CSV_MAX_FIELD_CHARACTERS)
NUMBER = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\Z")
IDENTIFIER = re.compile(r"(?:^|_)(?:id|code|sku|serial|barcode|zip|postcode)(?:$|_)|编号|编码|序列号|条码|邮编", re.I)
MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CHART = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
DRAW = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def source_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {"durationSeconds": value.total_seconds(), "sourceType": "timedelta"}
    if hasattr(value, 'text') and hasattr(value, 'ref'):
        return {"formulaText":value.text,"range":value.ref,"sourceType":type(value).__name__}
    return value


def _names(headers):
    display = [str(value) if value is not None and str(value) else "" for value in headers]
    counts = Counter(display)
    names = []
    for index, value in enumerate(display, 1):
        name = value if value and counts[value] == 1 else f"{value or 'unnamed'} [column {index}]"
        while name in names:
            name += " [source]"
        names.append(name)
    return names


def _numeric_column(name, values):
    present = [value for value in values if value != ""]
    # CSV has no native scalar types. Inference is explicit and conservative;
    # the exact strings remain authoritative in sourceValues.
    if not present or IDENTIFIER.search(name):
        return False
    for value in present:
        token = value.strip()
        if not NUMBER.fullmatch(token) or re.match(r"[+-]?0\d", token):
            return False
        if not math.isfinite(float(token)):
            return False
    return True


def parse_csv(text, units, *, delimiter=',', line_offset=0):
    reader = csv.reader(io.StringIO(text, newline=""), strict=True,delimiter=delimiter)
    headers, source_rows, blank_lines = None, [], []
    previous = line_offset
    for values in reader:
        start, end = previous + 1, reader.line_num+line_offset
        previous = end
        if not values:
            blank_lines.append(start)
            continue
        if headers is None:
            headers, header_line = values, start
            continue
        if len(values) > len(headers):
            raise ValueError(f"CSV_EXTRA_FIELDS_AT_LINE_{start}: header={len(headers)} row={len(values)}")
        source_rows.append((start, end, values))
    headers = headers or []
    names = _names(headers)
    numeric = {name for index, name in enumerate(names)
               if _numeric_column(name, [row[index] for _, _, row in source_rows if index < len(row)])}
    records = []
    for start, end, cells in source_rows:
        original = {name: cells[index] for index, name in enumerate(names) if index < len(cells)}
        values, numbers = {}, {}
        for name, value in original.items():
            if value == "":
                values[name] = None
            elif name in numeric:
                token = value.strip()
                values[name] = int(token) if re.fullmatch(r"[+-]?\d+", token) else float(token)
                numbers[name] = {"lexeme": value, "decimalValue": str(Decimal(token)), "sourceType": "csv-field"}
            else:
                values[name] = value
        records.append({"rawRow": start, "values": values, "sourceValues": original,
                        "numberLexemes": numbers,
                        "locator": {"lineStart": start, "lineEnd": end},
                        "valueTransforms": [{"field": name, "operation": "infer-csv-number"} for name in numeric if name in original]})
    payload = records_payload(records)
    # Header-only files still expose their declared fields.
    columns = {column['name']: column for column in payload['columns']}
    payload['columns'] = [{**columns.get(name, {"name": name, "dtype": "unknown", "nullable": True, "sample": []}),
                           "sourceHeader": headers[index], "sourceColumn": index + 1,
                           "typeProvenance": "conservative-lexical-inference" if name in numeric else "source-string"}
                          for index, name in enumerate(names)]
    payload['nullSemantics'] = "Empty CSV field has unspecified null-versus-empty-string meaning; sourceValues retains the empty string. Missing trailing fields remain absent."
    item = ParsedItem('table', payload, {"sheet": "root", "rawRowStart": records[0]['rawRow'] if records else previous + 1,
                      "headerLine": header_line if headers else None, "rowNumbering": "physical-line"}, units=dict(units))
    return [item], {"sheetCount": 1, "recipeVersion": RECIPE, "physicalLineCount": len(text.splitlines()),
                    "blankRecordLines": blank_lines, "nativeRecordsComplete": True,'sourceDelimiter':delimiter}, []


def _package_metadata(path):
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if any(entry.file_size > 128 * 1024 * 1024 for entry in entries) or sum(entry.file_size for entry in entries) > 512 * 1024 * 1024:
            raise ValueError("XLSX_UNCOMPRESSED_PACKAGE_LIMIT")
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        rels = {element.attrib['Id']: element for element in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        sheets = {}
        for sheet in workbook.findall(MAIN+'sheets/'+MAIN+'sheet'):
            relation = rels[sheet.attrib[REL+'id']]
            if relation.attrib.get('TargetMode') == 'External':
                raise ValueError('XLSX_EXTERNAL_WORKSHEET_UNSUPPORTED')
            target = relation.attrib['Target']
            member = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
            if not member.startswith('xl/') or '..' in member.split('/'):
                raise ValueError('XLSX_INVALID_WORKSHEET_MEMBER')
            xml = ET.fromstring(archive.read(member))
            cells = {}
            for cell in xml.iter(MAIN+'c'):
                formula = cell.find(MAIN+'f');value = cell.find(MAIN+'v')
                cells[cell.attrib['r']] = {"ooxmlType": cell.attrib.get('t', 'n'),
                    "rawValue": value.text if value is not None else None,
                    "formulaXml": formula.text if formula is not None else None,
                    "formulaAttributes": dict(formula.attrib) if formula is not None else {},
                    "hasFormula": formula is not None}
            sheets[sheet.attrib['name']] = {"cells": cells, "xmlMember": member}
        charts = []
        for entry in entries:
            if not re.fullmatch(r'xl/charts/chart\d+\.xml', entry.filename):
                continue
            xml = ET.fromstring(archive.read(entry))
            series = []
            for number, element in enumerate(xml.iter(CHART+'ser'), 1):
                refs = []
                for ref in element.iter():
                    if ref.tag not in {CHART+'numRef', CHART+'strRef'}:
                        continue
                    formula = ref.find(CHART+'f')
                    refs.append({"kind": ref.tag.rsplit('}',1)[-1],
                                 "reference": formula.text if formula is not None else None,
                                 "cachedPoints": [{"index": point.attrib.get('idx'),
                                     "value": point.findtext(CHART+'v')} for point in ref.iter(CHART+'pt')],
                                 "cacheTrust": "unverified"})
                series.append({"seriesIndex": number, "references": refs})
            charts.append({"chartId": entry.filename, "title": ''.join(node.text or '' for title in xml.iter(CHART+'title') for node in title.iter(DRAW+'t')),
                           "chartTypes": [element.tag.rsplit('}',1)[-1] for element in xml.iter() if element.tag.startswith(CHART) and element.tag.endswith('Chart')],
                           "series": series})
        return sheets, charts, any(entry.filename.startswith('xl/externalLinks/') for entry in entries)


def _header_row(sheet, width):
    rows = list(sheet.iter_rows(min_row=1, max_row=min(30, sheet.max_row), max_col=width))
    for cells in rows:
        values = [cell.value for cell in cells if cell.value is not None]
        if len(values) >= max(2, math.ceil(width * .6)) and all(isinstance(value, str) and value.strip() for value in values) and len(set(values)) == len(values):
            return cells[0].row
    return next((cells[0].row for cells in rows if any(cell.value is not None for cell in cells)), 1)


def parse_xlsx(path, units):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, range_boundaries
    metadata, charts, external = _package_metadata(path)
    source = load_workbook(path, data_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, keep_links=False)
    items, sheets, flags = [], [], []
    try:
        for sheet in source.worksheets:
            native = metadata[sheet.title]['cells']
            width = max((cell.column for row in sheet for cell in row if cell.value is not None), default=1)
            height = max((cell.row for row in sheet for cell in row if cell.value is not None), default=1)
            header = _header_row(sheet, width)
            sheet_flags = []
            from .source_timezone import worksheet_policy
            time_policy=worksheet_policy(sheet,header)
            sheet_flags.extend(time_policy['qualityFlags'])
            raw_headers = [source_value(sheet.cell(header, column).value) for column in range(1, width+1)]
            names = _names(raw_headers)
            records, cell_manifest = [], {}
            for row in sheet.iter_rows(min_row=1, max_row=height, max_col=width):
                for cell in row:
                    xml = native.get(cell.coordinate)
                    if xml is None and cell.value is None:
                        continue
                    xml = xml or {}
                    formula = cell.data_type == 'f' or xml.get('hasFormula', False)
                    value = cached[sheet.title][cell.coordinate].value if formula else cell.value
                    cache_status = ('missing' if xml.get('rawValue') is None else 'present-unverified') if formula else 'not-applicable'
                    cell_manifest[cell.coordinate] = {**xml, "cell": cell.coordinate,
                        "value": source_value(value), "sourceType": cell.data_type,
                        "numberFormat": cell.number_format, "isDate": getattr(cell,'is_date',False),
                        "formula": source_value(cell.value) if formula else None,
                        "cacheStatus": cache_status, "calculationEligible": not formula and cell.data_type!='e',
                        "hiddenRow": bool(sheet.row_dimensions[cell.row].hidden)}
                    if formula:
                        sheet_flags.append('FORMULA_CACHE_UNVERIFIED' if cache_status=='present-unverified' else 'FORMULA_CACHE_MISSING')
            for row_number in range(header+1, height+1):
                cells = [sheet.cell(row_number, column) for column in range(1,width+1)]
                if all(cell.value is None for cell in cells):
                    continue
                values, source_values, number_lexemes, formulas, eligibility = {}, {}, {}, {}, {}
                cell_refs = {}
                for name, cell in zip(names, cells):
                    info = cell_manifest.get(cell.coordinate, {})
                    values[name] = info.get('value')
                    source_values[name] = source_value(cell.value)
                    cell_refs[name] = cell.coordinate
                    eligibility[name] = info.get('calculationEligible', True)
                    if info.get('formula') is not None:
                        formulas[name] = {key:info.get(key) for key in ['formula','cacheStatus','rawValue','value','calculationEligible']}
                    elif info.get('ooxmlType')=='n' and info.get('rawValue') is not None and not getattr(cell,'is_date',False):
                        number_lexemes[name] = {"lexeme": info['rawValue'], "cell": cell.coordinate, "sourceType": "ooxml-number"}
                records.append({"rawRow": row_number, "values": values, "sourceValues": source_values,
                    "numberLexemes": number_lexemes, "formulas": formulas, "calculationEligibility": eligibility,
                    "locator": {"sheet": sheet.title,"range": f"A{row_number}:{get_column_letter(width)}{row_number}","cells": cell_refs}})
            payload = records_payload(records)
            payload['sourceTimePolicy']=time_policy
            columns = {column['name']:column for column in payload['columns']}
            payload['columns'] = [{**columns.get(name,{'name':name,'dtype':'unknown','nullable':True,'sample':[]}),
                "sourceHeader": raw_headers[index], "headerCell": f"{get_column_letter(index+1)}{header}",
                "sourceColumn": index+1} for index,name in enumerate(names)]
            payload['coverage'].update(headerRow=header, lastPhysicalRow=height, includesHiddenRows=True, includesHiddenColumns=True)
            items.append(ParsedItem('table',payload,{"sheet":sheet.title,"headerRow":header,"rawRowStart":header+1,
                "range":f"A{header}:{get_column_letter(width)}{height}"}, units=dict(units), quality_flags=sorted(set(sheet_flags))))
            # A bad declaration on one worksheet does not taint a different
            # worksheet's explicit time basis. Per-sheet flags remain native.
            flags.extend(flag for flag in sheet_flags if not flag.startswith('SOURCE_TIMEZONE_'))
            sheets.append({"sheet":sheet.title,"sheetState":sheet.sheet_state,"headerRow":header,
                "sourceTimePolicy":time_policy,
                "range":f"A1:{get_column_letter(width)}{height}","cells":cell_manifest,
                "mergedRanges":[str(value) for value in sheet.merged_cells.ranges],
                "hiddenRows":[index for index,dim in sheet.row_dimensions.items() if dim.hidden],
                "hiddenColumns":[{"min":dim.min,"max":dim.max,"label":key} for key,dim in sheet.column_dimensions.items() if dim.hidden]})
        by_sheet={sheet['sheet']:sheet for sheet in sheets}
        for chart in charts:
            for series in chart['series']:
                for ref in series['references']:
                    expression=ref['reference'] or ''
                    if '!' not in expression or '[' in expression:
                        ref['readStatus']='UNSUPPORTED_REFERENCE';continue
                    sheet_name,area=expression.rsplit('!',1);sheet_name=sheet_name.strip("'").replace("''", "'")
                    if sheet_name not in by_sheet:
                        ref['readStatus']='SOURCE_SHEET_NOT_FOUND';continue
                    try:
                        left,top,right,bottom=range_boundaries(area)
                        if None in (left,top,right,bottom) or (right-left+1)*(bottom-top+1)>50000:
                            raise ValueError('CHART_RANGE_LIMIT')
                        ref['sourceCells']=[{"sheet":sheet_name,"cell":f'{get_column_letter(column)}{row}',
                            **by_sheet[sheet_name]['cells'].get(f'{get_column_letter(column)}{row}',{'value':None})}
                            for row in range(top,bottom+1) for column in range(left,right+1)]
                        ref['readStatus']='READ_NATIVE_CELLS'
                    except (ValueError,TypeError):ref['readStatus']='UNSUPPORTED_OR_OVERSIZED_RANGE'
        calculation=source.calculation
        attributes={"sheetCount":len(sheets),"recipeVersion":RECIPE,"nativeCellsComplete":True,
            "workbook":{"dateSystem":"1904" if source.epoch.year==1904 else "1900", "epoch":source.epoch.isoformat(),
                "sheets":sheets,"charts":charts,"externalLinksPresent":external,
                "calculationProperties":dict(calculation) if calculation is not None else {},
                "formulaExecution":"none","externalLinkUpdate":"disabled","cacheTrust":"unverified"}}
        return items, attributes, sorted(set(flags))
    finally:
        source.close();cached.close()
